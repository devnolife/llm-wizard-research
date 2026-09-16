#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Bangun deck sidang tesis (PPTX, 16:9, bisa diedit) dari isi BAB I–V.

Jalankan:  python3 drafts/sidang/build_deck.py
Keluaran:  drafts/sidang/SIDANG_TESIS_Wizard_Research.pptx

Setiap slide memuat catatan pembicara (waktu, poin kunci, transisi).
Slide cadangan (setelah "Terima kasih") disiapkan untuk sesi tanya-jawab.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
FIG = HERE / "figures"
PAPER_FIG = ROOT / "paper_figures"
OUT = HERE / "SIDANG_TESIS_Wizard_Research.pptx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
RED = RGBColor(0xB2, 0x22, 0x22)
GREY = RGBColor(0x55, 0x5F, 0x6B)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xC7, 0x77, 0x00)
BLACK = RGBColor(0x11, 0x11, 0x11)

W, H = Inches(13.333), Inches(7.5)
FOOTER = "Sidang Tesis — Andi Agung Dwi Arya B (D082251054) · Magister Teknik Informatika, Universitas Hasanuddin"

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
_counter = {"n": 0}


# ----------------------------------------------------------------- helpers
def _txt(shape, text, size=20, bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return tf


def new_slide(title, subtitle=None, section=None):
    _counter["n"] += 1
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.05))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.05), W, Inches(0.06))
    acc.fill.solid(); acc.fill.fore_color.rgb = RED; acc.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.45), Inches(0.08), Inches(10.9), Inches(0.92))
    tsize = 28 if len(title) <= 52 else (24 if len(title) <= 66 else 21)
    _txt(tb, title, size=tsize, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    if section:
        sb = s.shapes.add_textbox(Inches(11.4), Inches(0.2), Inches(1.8), Inches(0.7))
        _txt(sb, section, size=12, color=RGBColor(0xCF, 0xD8, 0xE3), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        st = s.shapes.add_textbox(Inches(0.45), Inches(1.18), Inches(12.4), Inches(0.5))
        _txt(st, subtitle, size=16, color=GREY)
    ft = s.shapes.add_textbox(Inches(0.45), Inches(7.05), Inches(11.5), Inches(0.35))
    _txt(ft, FOOTER, size=10, color=GREY)
    num = s.shapes.add_textbox(Inches(12.3), Inches(7.05), Inches(0.8), Inches(0.35))
    _txt(num, str(_counter["n"]), size=11, color=GREY, align=PP_ALIGN.RIGHT)
    return s


def bullets(slide, items, left=0.5, top=1.75, width=12.3, height=5.1, size=20, color=BLACK):
    """items: str | (str, level) | (str, level, color)."""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    first = True
    for it in items:
        text, level, col = (it, 0, color) if isinstance(it, str) else (it + (color,))[:3]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        prefix = "• " if level == 0 else "– "
        p.level = min(level, 2)
        p.space_after = Pt(6 if level == 0 else 2)
        r = p.add_run()
        r.text = prefix + text
        r.font.size = Pt(size - 3 * level)
        r.font.color.rgb = col
        r.font.name = "Calibri"
    return tb


def callout(slide, text, left, top, width, height=0.8, fill=NAVY, color=WHITE, size=17, bold=True):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid(); box.fill.fore_color.rgb = fill; box.line.fill.background()
    box.adjustments[0] = 0.12
    _txt(box, text, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return box


def picture(slide, path, left, top, width=None, height=None, center=False):
    kw = {}
    if width: kw["width"] = Inches(width)
    if height: kw["height"] = Inches(height)
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kw)
    if center:
        pic.left = int((W - pic.width) / 2)
    return pic


def table(slide, data, left, top, width, col_w=None, size=13, header_fill=NAVY, row_h=0.38):
    rows, cols = len(data), len(data[0])
    shp = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(row_h * rows))
    t = shp.table
    if col_w:
        for i, w in enumerate(col_w):
            t.columns[i].width = Inches(w)
    for r in range(rows):
        for c in range(cols):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            val = data[r][c]
            color, bold = BLACK, False
            if isinstance(val, tuple):
                val, color = val[0], val[1]
                bold = True
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = str(val)
            run.font.size = Pt(size)
            run.font.name = "Calibri"
            if r == 0:
                run.font.bold = True
                run.font.color.rgb = WHITE
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
            else:
                run.font.color.rgb = color
                run.font.bold = bold
                cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT if r % 2 else WHITE
    return shp


def notes(slide, minutes, points, transition=None):
    body = [f"[Waktu: {minutes}]", ""]
    body += [f"• {p}" for p in points]
    if transition:
        body += ["", f"→ Transisi: {transition}"]
    slide.notes_slide.notes_text_frame.text = "\n".join(body)


# ================================================================== SLIDES
# 1. Judul
s = prs.slides.add_slide(BLANK); _counter["n"] += 1
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.55), Inches(2.2), Inches(0.08)); acc.fill.solid(); acc.fill.fore_color.rgb = RED; acc.line.fill.background()
_txt(s.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.8), Inches(0.5)), "SIDANG TESIS MAGISTER", size=16, color=RGBColor(0xCF, 0xD8, 0xE3))
_txt(s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.8), Inches(2.1)),
     ["Pendekatan Neuro-Symbolic Agentic untuk Deteksi", "Indikator Synthesis Gap pada Literatur Ilmiah"], size=36, bold=True, color=WHITE)
_txt(s.shapes.add_textbox(Inches(0.8), Inches(3.75), Inches(11.8), Inches(0.9)),
     "Wizard Research — sistem bantu keputusan yang menggabungkan penalaran LLM dengan validasi berbasis aturan", size=18, color=RGBColor(0xE5, 0xE7, 0xEB))
_txt(s.shapes.add_textbox(Inches(0.8), Inches(4.9), Inches(7), Inches(1.6)),
     ["Andi Agung Dwi Arya B — D082251054", "Program Studi Magister Teknik Informatika", "Fakultas Teknik, Universitas Hasanuddin"], size=18, color=WHITE)
_txt(s.shapes.add_textbox(Inches(8.2), Inches(4.9), Inches(4.6), Inches(1.6)),
     ["Pembimbing: [nama pembimbing 1]", "[nama pembimbing 2]", "", "Makassar, [tanggal sidang] 2026"], size=15, color=RGBColor(0xCF, 0xD8, 0xE3))
notes(s, "0:30", ["Salam, perkenalan singkat, sebutkan judul lengkap.",
                  "Satu kalimat posisi: ini alat bantu keputusan, bukan pengganti peneliti — tema yang akan berulang.",
                  "Isi nama pembimbing dan tanggal sebelum sidang."],
      "Mulai dari fenomena: mengapa mencari research gap itu sulit.")

# 2. Latar belakang
s = new_slide("Latar Belakang: Menemukan Research Gap Adalah Tugas Sintesis", section="PENDAHULUAN")
bullets(s, [
    "Peneliti membaca banyak literatur, membandingkan temuan, menalar induktif, lalu menyimpulkan apa yang belum terjawab secara kolektif (Cooper, 1998; Booth et al., 2012)",
    "Pipeline RAG+LLM linear (retrieve-then-generate) bekerja pada level kesamaan vektor & generasi probabilistik:",
    ("tanpa penalaran bertahap, tanpa validasi logis, tidak membedakan asosiasi semantik dari hubungan kausal", 1),
    ("rawan korelasi semu dan over-claiming (Bender & Koller, 2020; Marcus & Davis, 2020)", 1),
    "Paradigma agen (multi-step reasoning + rule-based validation) secara struktural lebih dekat ke proses kognitif peneliti — tetapi belum diketahui sejauh mana ia mampu, dan di mana batasnya",
], size=20, height=3.6)
callout(s, "Kesenjangan: belum ada bukti empiris tentang kemampuan — dan batas — pendekatan agentic + validasi simbolis untuk mendeteksi indikator synthesis gap",
        0.5, 5.6, 12.3, 1.0, fill=RED, size=17)
notes(s, "1:00", ["Tekankan pola FENOMENA → KESENJANGAN (respons kritik #1 penguji: rumusan masalah harus berupa masalah, bukan spesifikasi teknis).",
                  "Fenomena: sintesis literatur adalah penalaran induktif; alat yang ada beroperasi di level semantik.",
                  "Kesenjangan: belum ada bukti empiris sejauh mana pendekatan agentic+simbolis mampu."],
      "Kesenjangan itu diturunkan menjadi tiga pertanyaan penelitian.")

# 3. Rumusan masalah
s = new_slide("Rumusan Masalah dan Tujuan", section="PENDAHULUAN")
table(s, [
    ["", "Pertanyaan Penelitian", "Tujuan"],
    ["RQ1", "Sejauh mana pendekatan agentic multi-step reasoning + rule-based validation mampu mendeteksi indikator synthesis gap (fragmentasi, inkonsistensi, ketidaklengkapan kolektif, ketiadaan dukungan bukti)?", "Mengukur kemampuan deteksi empat indikator"],
    ["RQ2", "Bagaimana mekanisme pembeda asosiasi semantik–hubungan logis dan rule-based validation memengaruhi akurasi dan false discovery rate?", "Mengisolasi kontribusi tiap komponen (ablasi)"],
    ["RQ3", "Apa batasan epistemologis pendekatan ini dibandingkan penalaran logis-induktif peneliti manusia?", "Mendeskripsikan batas sistem secara eksplisit"],
], 0.5, 1.75, 12.3, col_w=[0.8, 8.0, 3.5], size=15, row_h=0.9)
callout(s, "Definisi operasional (Cooper 1998; Booth 2012): synthesis gap = fragmentasi · inkonsistensi · ketidaklengkapan kolektif · + ketiadaan dukungan bukti (indikator ke-4, kontribusi penelitian ini)",
        0.5, 5.65, 12.3, 1.0, fill=NAVY, size=15)
notes(s, "1:15", ["Baca RQ dengan kata kunci 'Sejauh mana' — bukan 'Bagaimana merancang' (revisi atas kritik penguji).",
                  "RQ3 sengaja ada: mengakui batas kemampuan sistem adalah kontribusi ilmiah, bukan kelemahan.",
                  "Definisi synthesis gap mengikuti literatur mainstream (respons kritik #2)."],
      "Untuk menjawabnya kami membangun sistem dengan empat fase.")

# 4. Kebaruan
s = new_slide("Kebaruan: Dari Pipeline Linear ke Neuro-Symbolic Agentic", section="PENDAHULUAN")
table(s, [
    ["Aspek", "Pipeline RAG+LLM biasa", "Wizard Research"],
    ["Alur", "retrieve → generate (satu langkah)", "Observe → Think → Act → Evaluate (LangGraph, maks. 3 iterasi, self-critique)"],
    ["Representasi pengetahuan", "embedding saja", "embedding + Fact Table SPO (8 tipe entitas, 12 predikat + 2 turunan)"],
    ["Validasi keluaran", "tidak ada", "Rule Engine 9 aturan (F1–F3, C1–C3, K1–K3) → PASS / FLAG / REJECT"],
    ["Semantik vs logis", "tidak dibedakan", "3 lapis: penanda linguistik → analisis struktural → NLI cross-encoder"],
    ["Keyakinan keluaran", "skor mentah LLM, satu run", "kalibrasi + abstensi selektif + rantai provenans + pelaporan k/n"],
    ["Klaim epistemologis", "\"research gaps\"", "\"indikator gap\" yang wajib divalidasi manusia"],
], 0.5, 1.7, 12.3, col_w=[2.6, 3.9, 5.8], size=14, row_h=0.62)
notes(s, "1:00", ["Ini respons langsung atas kritik #4 (RAG+LLM bukan hal baru): kebaruan digeser ke integrasi neural + simbolik, bukan sekadar pipeline.",
                  "Baris 'validasi' dan 'semantik vs logis' menjawab kritik #5 dan #6.",
                  "Baris terakhir: posisi epistemologis — kita hanya mengklaim indikator."],
      "Mari lihat arsitekturnya.")

# 5. Arsitektur
s = new_slide("Arsitektur Empat Fase", section="METODE")
picture(s, PAPER_FIG / "fig1_architecture.png", 0.6, 1.4, width=8.3)
bullets(s, [
    "Fase 1 — Ingestion: PDF → pembersihan → seksi IMRaD → chunk sadar-kalimat → ChromaDB (embedder multibahasa + reranker)",
    "Fase 2 — Fact Extraction: LLM (JSON mode + retry + salvage) + pola linguistik → triple SPO",
    "Fase 3 — Agentic Analysis: koordinator LangGraph dengan 5 tool (RAG, PaperAnalyzer, NLIChecker, KGQuerier, SelfCritic)",
    "Fase 4 — Logical Validation: Rule Engine atas fakta KG; verdict difusikan ke kalibrasi",
], left=9.0, top=1.5, width=4.1, height=5.2, size=14)
notes(s, "1:15", ["Jalankan satu 'permintaan' dari kiri ke kanan: PDF masuk, fakta diekstrak, agen berputar Observe-Think-Act-Evaluate, Rule Engine memvalidasi.",
                  "Tekankan Fase 4 sebagai komponen baru hasil revisi (kritik #7: diagram jangan black box)."],
      "Dua komponen simbolik inti: Fact Table dan Rule Engine.")

# 6. Komponen simbolik
s = new_slide("Komponen Simbolik: Fact Table SPO dan Rule Engine 9 Aturan", section="METODE")
table(s, [
    ["Kategori", "Aturan", "Pertanyaan yang diuji", "Verdict bila gagal"],
    ["Kelayakan", "F1 sumber daya · F2 data · F3 skala", "Apakah metode layak untuk domain/kendala yang diklaim?", "REJECT (F1, F3) / FLAG (F2)"],
    ["Kausalitas", "C1 bukti kausal · C2 arah · C3 confounding", "Apakah klaim kausal punya bukti; adakah jalur alternatif di KG?", "FLAG"],
    ["Konsistensi", "K1 kontradiksi internal · K2 konsistensi fakta KG · K3 transitivitas", "Apakah klaim bertentangan dengan fakta lain di Fact Table?", "FLAG"],
], 0.5, 1.7, 12.3, col_w=[1.7, 3.6, 4.6, 2.4], size=14, row_h=0.7)
bullets(s, [
    "Fact Table: (Paper_A, USES_METHOD, CNN) · (CNN, APPLIES_TO, Medical_Imaging) · (ResNet, IMPROVES, VGG) — 8 tipe entitas, 12 predikat dasar + 2 turunan (INFEASIBLE_FOR, CORRELATES_WITH)",
    "Verdict = bukti tentang klaim, bukan pendapat kedua: PASS menguatkan (×1,10), FLAG mendiskon (×0,80), REJECT membatalkan",
    "Validasi berjalan < 0,01 detik — lapisan logis praktis tanpa overhead",
], top=4.75, height=2.2, size=16)
notes(s, "1:30", ["Respons kritik #8 (KG tanpa tabel fakta): KG kini Fact Base dengan ontologi eksplisit.",
                  "Contoh aturan konkret: F1 — metode butuh GPU besar tapi domain edge device → REJECT.",
                  "Jelaskan bahwa verdict masuk ke kalibrasi, bukan dirata-ratakan."],
      "Bagaimana sistem membedakan 'sering muncul bersama' dari 'benar-benar berhubungan'?")

# 7. Semantik vs logis + NLI
s = new_slide("Membedakan Asosiasi Semantik dari Hubungan Logis", section="METODE")
bullets(s, [
    "Lapisan 1 — Penyaringan semantik: kandidat pasangan klaim dari kesamaan embedding & ko-okurensi",
    "Lapisan 2 — Ekstraksi bukti: penanda linguistik (outperforms, contradicts, is based on) + analisis struktural seksi",
    "Lapisan 3 — Verifikasi: NLI cross-encoder terdedikasi (entailment / neutral / contradiction) + fakta KG",
    "Adjudikasi kontradiksi vs heterogenitas: klaim dinormalisasi (arah efek, PICO) lalu diuji Cochran Q, I², τ² — beda hasil karena desain berbeda ≠ kontradiksi",
    "Indikator yang hanya berasal dari asosiasi semantik selalu berlabel requires_human_validation",
], size=18, height=3.7)
callout(s, "Ditambah lapisan bukti bebas-LLM: kopling bibliografis (Kessler, 1963) dari daftar pustaka yang diurai, dan korroborasi pernyataan keterbatasan/future work penulis — sebagai BUKTI, bukan skor",
        0.5, 5.55, 12.3, 1.1, fill=NAVY, size=15)
notes(s, "1:15", ["Ini jawaban desain untuk RQ2 dan kritik #5.",
                  "Sebut bahwa NLI adalah sinyal independen yang nanti diuji ablasi (H9).",
                  "Kopling bibliografis dan korroborasi penulis: tambahan pasca-revisi, keduanya tidak bergantung pada LLM."],
      "Terakhir dari metode: bagaimana keyakinan dilaporkan secara jujur.")

# 8. Kalibrasi, provenans, k/n
s = new_slide("Keyakinan yang Dapat Dipertanggungjawabkan", section="METODE")
table(s, [
    ["Mekanisme", "Aturan", "Tujuan"],
    ["Rantai provenans", "klaim → jurnal terkutip → kutipan verbatim → verdict; mata rantai putus → ditahan", "Gap yang tak bisa ditelusuri ke kutipan bukan temuan"],
    ["Kalibrasi post-hoc", "temperature scaling + ambang konformal (α = 0,10); identitas sampai ≥ 4 label pakar", "Keyakinan 90% harus benar ≈ 90%"],
    ["Abstensi selektif", "terkalibrasi < 0,45 atau di bawah ambang konformal → needs_review", "Menahan, bukan menyajikan, klaim lemah"],
    ["Pelaporan k/n", "n = 3 run identik; stabil bila muncul ≥ ⌈2n/3⌉ = 2; k/n = anotasi & filter, bukan skor", "Satu run LLM = satu undian"],
    ["Kebaruan usulan", "band derivative / sweet spot (0,25–0,65) / off-topic; hanya untuk pemeringkatan", "Usulan tak berjangkar pada indikator turun prioritas"],
], 0.5, 1.7, 12.3, col_w=[2.3, 6.2, 3.8], size=14, row_h=0.72)
notes(s, "0:45", ["Tiga hal ini yang membedakan 'indikator' dari 'klaim': bisa ditelusuri, terkalibrasi, dan dilaporkan dengan sebaran.",
                  "Tekankan: kalibrator sengaja identitas sebelum ada label pakar dan status itu ditampilkan di antarmuka."],
      "Sekarang desain eksperimennya.")

# 9. Desain eksperimen
s = new_slide("Desain Eksperimen: Dua Korpus, Dua Rezim, 14 Metrik", section="EVALUASI")
table(s, [
    ["", "Benchmark & ablasi (4.3.1–4.3.7)", "Korpus aplikatif (4.3.8–4.3.11)"],
    ["Korpus", "23 paper CS (T1–T4: arsitektur DL, CV, NLP, edge), Inggris", "35 jurnal forensika digital, Indonesia + Inggris"],
    ["LLM", "llama3.2 3B & gpt-oss 13B (Ollama, lokal, seed)", "claude-opus-4.8-fast (Copilot SDK)"],
    ["Tujuan", "H6, H7, H9, H10 multi-run; adversarial; kontrol negatif", "sistem lengkap pada dokumen nyata: M9–M14"],
    ["Mode", "full · no-rule-engine · linear-baseline · nli/no-nli · cross-critic (+TC)", "pipeline penuh, 2 eksekusi versi akhir"],
], 0.5, 1.7, 12.3, col_w=[1.5, 5.6, 5.2], size=14, row_h=0.6)
bullets(s, [
    "Metrik M1–M8 (jumlah, distribusi, confidence, verdict, waktu, validasi manusia, RERR, akurasi adversarial) + M9–M14 (kalibrasi, abstensi, provenans, kebaruan, stabilitas k/n, korroborasi)",
    "Statistik: satu observasi per run, Mann–Whitney U, koreksi Holm–Bonferroni (keluarga 8 uji), Cliff's δ, CI bootstrap",
    "Metrik pakar (EAR, LCS, AS, FDR, SHG, REP) melalui instrumen penilaian blinded, Cohen's κ ≥ 0,6",
], top=4.85, height=2.1, size=15)
notes(s, "1:15", ["Jelaskan mengapa dua rezim dan mengapa TIDAK dibandingkan silang: model lokal untuk seed/ulangan; model penyedia untuk bahasa Indonesia.",
                  "Sebut angka: 14 metrik, 10 hipotesis, 1 kontrol negatif."],
      "Hasil pertama: apakah keempat indikator terdeteksi dan apakah Rule Engine benar-benar menyaring.")

# 10. Hasil benchmark + adversarial
s = new_slide("Hasil Benchmark: Deteksi Indikator dan Bukti Rule Engine Diskriminatif", section="HASIL")
table(s, [
    ["Topik", "Indikator", "Frag.", "Inkons.", "Ketidaklengkapan", "Confidence"],
    ["T1 Deep learning", "3", "1", "1", "1", "0,717"],
    ["T2 Computer vision", "4", "1", "1", "2", "0,688"],
    ["T3 NLP & attention", "4", "1", "1", "2", "0,688"],
    ["T4 Edge/efisien", "3", "1", "1", "1", "0,717"],
    [("Total", NAVY), ("14", NAVY), "4", "4", "6", ("0,700", NAVY)],
], 0.5, 1.7, 6.3, col_w=[2.1, 0.9, 0.7, 0.8, 1.1, 0.7], size=13, row_h=0.42)
bullets(s, [
    "248 fakta SPO dari 23 paper (22/23 ≥ 1 fakta); 14 indikator, 100% PASS pada data bersih (RERR 0%)",
    "Multi-run (5 run): 17,2 ± 3,2 indikator; RERR 3,8 ± 8,5% — angka satu run adalah satu realisasi",
    "Inkonsistensi terlemah (0,50): selalu requires_human_validation",
], left=0.5, top=4.4, width=6.3, height=2.5, size=14)
picture(s, PAPER_FIG / "fig4_adversarial.png", 7.0, 1.55, width=5.9)
callout(s, "Validasi adversarial 6/6 benar: F1/F3 → REJECT (−0,60), F2/K1/C1 → FLAG (−0,15…−0,20), kontrol → PASS. PASS 100% bukan karena engine meloloskan semua.",
        7.0, 5.6, 5.9, 1.25, fill=NAVY, size=13)
notes(s, "1:15", ["Angka utama: 14 indikator, 248 fakta, 100% PASS — lalu segera jawab keraguan yang pasti muncul: 'apakah Rule Engine bekerja?' dengan adversarial 6/6.",
                  "Sebut bahwa PASS 100% bukan sifat tetap (multi-run RERR 3,8%)."],
      "Komponen mana yang benar-benar berkontribusi? Ablasi multi-run.")

# 11. Ablasi
s = new_slide("Ablasi Multi-Run: NLI Signifikan, Rule Engine Kualitatif, Kritikus Memangkas", section="HASIL")
picture(s, FIG / "modes_multirun.png", 0.5, 1.4, height=4.3, center=True)
callout(s, "H9 terkonfirmasi · H6/H7 tidak signifikan pada proksi kuantitatif (efeknya akuntabilitas & penolakan adversarial) · H10 terkonfirmasi arah: 'sedikit tetapi yakin'",
        0.5, 6.05, 12.3, 0.85, fill=NAVY, size=14)
notes(s, "1:30", ["Baca grafik dari atas: baseline linear menghasilkan LEBIH BANYAK indikator (20) tetapi tanpa fakta, verdict, dan jejak — dan identik antar topik (template).",
                  "H9: NLI menambah 8 indikator/run dan menaikkan confidence, p Holm 0,043 setelah koreksi 8 uji.",
                  "H6/H7 tidak signifikan: jujur katakan power kecil (5 run) dan efek Rule Engine bukan kuantitas.",
                  "H10: kritikus model berbeda menolak 85,6% kandidat."],
      "Kalau kritikus begitu agresif — apakah sistem masih bisa tertipu topik fiktif?")

# 12. Kontrol negatif
s = new_slide("Kontrol Negatif: Topik yang Tidak Ada di Korpus", section="HASIL")
picture(s, FIG / "crosscritic_control.png", 0.5, 1.4, height=4.45, center=True)
callout(s, "4/5 run: 0 indikator palsu. 1 run: satu indikator palsu lolos LLM utama + Rule Engine + kritikus dengan confidence 0,885 → hanya syarat kutipan verbatim (provenans) yang secara prinsip menolak klaim tentang topik yang tidak ada",
        0.5, 6.0, 12.3, 0.95, fill=RED, size=14)
notes(s, "1:15", ["Ini slide kejujuran: sistem TIDAK sempurna dan kami menemukannya sendiri.",
                  "Satu false gap berkeyakinan tinggi adalah argumen empiris terkuat untuk rantai provenans.",
                  "Kontrol negatif baru dijalankan pada mode cross-critic — keterbatasan yang dicatat."],
      "Sekarang sistem lengkap pada dokumen nyata berbahasa Indonesia.")

# 13. Korpus aplikatif
s = new_slide("Korpus Aplikatif 35 Jurnal Forensika Digital: Keempat Tipe Indikator Terdeteksi", section="HASIL")
picture(s, FIG / "forensik_indicators.png", 0.5, 1.35, height=4.05, center=True)
bullets(s, [
    "Kopling bibliografis (bebas LLM): 26 jurnal → 19 kelompok tanpa rujukan bersama; 318/325 pasangan terputus (isolasi 0,978); modularitas 0,70; 3 sitasi langsung terverifikasi",
    "Ketiadaan dukungan bukti: 3 klaim pembuka diasersikan lintas jurnal tanpa satu paragraf bukti primer (kemiripan terbaik 0,00 < 0,45); 1 citation echo",
    "Provenans lengkap 3/4; yang ke-4 ditahan — abstensi selektif membedakan, tidak menyala untuk semua",
], top=5.55, height=1.45, size=13)
notes(s, "1:45", ["Tiga sinyal independen menunjuk fragmentasi yang sama: kopling (daftar pustaka), isolasi (embedding), peta cakupan.",
                  "Indikator ke-4 (support gap) muncul pertama kali pada versi akhir; usulan yang berjangkar padanya naik ke prioritas high.",
                  "Korroborasi penulis 2/2 indikator yang dapat dikorroborasi — sebagai bukti, confidence tidak berubah.",
                  "Waktu ~12 menit untuk 35 jurnal; LLM Copilot."],
      "Seberapa stabil temuan yang bergantung pada LLM?")

# 14. Stabilitas k/n
s = new_slide("Stabilitas Lintas-Run (M13): Angka Satu Run Adalah Satu Undian", section="HASIL")
picture(s, FIG / "kn_stability.png", 0.5, 1.35, height=4.1, center=True)
bullets(s, [
    "Hanya pernyataan stabil (k ≥ 2) yang diteruskan ke korroborasi & pemeringkatan; yang k = 1 disimpan dan ditampilkan sebagai tidak stabil",
    "Kualitas isi terhadap gold standard eksternal (Mendeley, 30 paper): kemiripan semantik 0,603; LLM-judge 2,75/5 — menangkap topik gap, kurang spesifik dibanding kurator manusia",
], top=5.6, height=1.4, size=14)
notes(s, "1:00", ["Pesan: reproducibility diukur, bukan diasumsikan. 58% stabil; gap implisit paling rapuh (41%).",
                  "Ini juga menjelaskan mengapa semua temuan LLM dilaporkan dengan k/n."],
      "Sekarang pemetaan jujur: hipotesis mana yang terbukti, mana yang belum.")

# 15. Status hipotesis
s = new_slide("Status Hipotesis dan Kriteria Keberhasilan", section="HASIL")
table(s, [
    ["Hipotesis", "Bukti", "Status"],
    ["H9 NLI menambah deteksi", "Δmed = 8, p Holm 0,043, δ = 1,0", ("Terkonfirmasi", GREEN)],
    ["H10 cross-critic memangkas indikator palsu", "85,6% ditolak; Δmed = −19, p Holm 0,043; conf 0,791 → 0,884", ("Terkonfirmasi (arah)", GREEN)],
    ["H7 Rule Engine menurunkan FDR", "diskriminatif 6/6 adversarial; proksi tidak signifikan; FDR butuh pakar", ("Sebagian", AMBER)],
    ["H6 agentic > linear", "kualitatif jelas; proksi p Holm ≥ 0,48", ("Tidak signifikan", AMBER)],
    ["H3 KG → lebih didukung bukti", "248 fakta + trace vs 0; provenans 3/4; korroborasi 2/2", ("Kualitatif", AMBER)],
    ["H4 EAR ≥ 50% · H5 LCS ≥ 3,5", "instrumen siap, label pakar belum terkumpul", ("Tertunda", RED)],
    ["H1, H2, H8", "tak ada pembanding bermakna di sistem akhir; user study tidak dilaksanakan", ("Tidak diuji", RED)],
    ["Kontrol negatif", "0,2 ± 0,45 false gap/run (cross-critic); mode lain belum", ("Sebagian", AMBER)],
], 0.5, 1.6, 12.3, col_w=[3.6, 6.2, 2.5], size=13, row_h=0.5)
callout(s, "Keempat kriteria 3.7.5 (EAR ≥ 50%, LCS ≥ 3,5, FDR turun ≥ 20%, REP ≥ 70%) bergantung pada label pakar → belum dapat dinyatakan terpenuhi maupun gagal. Kriteria teknis lain terpenuhi seluruhnya.",
        0.5, 6.15, 12.3, 0.8, fill=NAVY, size=13)
notes(s, "1:30", ["Slide ini sengaja ada agar penguji melihat kami memetakan janji BAB III ke bukti BAB IV tanpa menyembunyikan apa pun.",
                  "Jelaskan: H1–H2 berasal dari proposal awal; setelah revisi arsitektur pertanyaannya dijawab H6/H7/H9."],
      "Dengan peta itu, jawaban atas tiga RQ.")

# 16. Kesimpulan RQ
s = new_slide("Kesimpulan: Jawaban atas Tiga Pertanyaan Penelitian", section="KESIMPULAN")
bullets(s, [
    ("RQ1 — Sejauh mana?", 0, NAVY),
    ("Keempat indikator terdeteksi dengan bukti yang dapat diaudit pada dua bahasa; sinyal bebas-LLM (kopling, kutipan) sebagai jangkar. Batas: 2–4 indikator/topik, 58% stabilitas k/n, inkonsistensi terlemah; Precision/Recall vs pakar belum terukur.", 1),
    ("RQ2 — Pengaruh mekanisme pembeda & validasi?", 0, NAVY),
    ("NLI menaikkan deteksi signifikan (H9); Rule Engine diskriminatif (6/6) dan verdict-nya mengatur kalibrasi & abstensi; kritikus lintas-model memangkas 85,6% kandidat (H10); false-gap 0,2/run. Besaran penurunan FDR menunggu label pakar.", 1),
    ("RQ3 — Batasan epistemologis?", 0, NAVY),
    ("Mendeteksi, tidak menilai kebermaknaan · deduksi terbatas pada fakta yang terekstrak · stokastisitas adalah sifat (k/n) · sinyal paling sederhana paling andal · tidak menjelaskan 'mengapa'. Posisi: decision support tool.", 1),
], size=18, height=5.2)
notes(s, "1:30", ["Jawab setiap RQ dalam dua kalimat: apa yang terbukti, apa batasnya.",
                  "Tutup RQ3 dengan posisi sistem — konsisten dengan slide pertama."],
      "Kontribusi yang bisa dibawa pulang.")

# 17. Kontribusi
s = new_slide("Kontribusi Penelitian", section="KESIMPULAN")
bullets(s, [
    "Rule-Based Validation Layer — 9 aturan, 3 verdict, independen dari LLM, < 0,01 s, diskriminatif (adversarial 6/6)",
    "Fact Table SPO — 8 tipe entitas, 12 + 2 predikat; grounding klaim pada fakta terstruktur",
    "Operasionalisasi empat indikator synthesis gap, termasuk indikator baru: ketiadaan dukungan bukti (uji kegagalan retrieval leave-one-out)",
    "Sinyal fragmentasi bebas-LLM via kopling bibliografis dari daftar pustaka yang diurai (isolasi 0,978 pada korpus forensik)",
    "Rantai provenans + kalibrasi + abstensi selektif sebagai satu kesatuan — terbukti diskriminatif (1/4 ditahan)",
    "Lapisan korroborasi pernyataan penulis sebagai bukti (bukan gap, bukan skor)",
    "Protokol pelaporan k/n untuk temuan bergantung-LLM, dengan bukti empirisnya (58,1% stabil; Jaccard 0,45–0,65)",
], size=17, height=5.2)
notes(s, "1:00", ["Tujuh kontribusi; tiga pertama adalah respons langsung revisi penguji, empat berikutnya lahir dari pengujian pada korpus nyata."],
      "Apa yang TIDAK kami klaim, dan apa langkah berikutnya.")

# 18. Batas klaim & saran
s = new_slide("Batas Klaim dan Saran", section="KESIMPULAN")
bullets(s, [
    ("Penelitian ini TIDAK menyimpulkan:", 0, RED),
    ("indikator = genuine gap (EAR/LCS/AS/FDR/SHG/REP belum diukur); ECE/Brier/AURC (kalibrator identitas); perbandingan silang dua rezim LLM; H1/H2/H8 teruji; false-gap sistem lengkap ≈ 0", 1),
], left=0.5, top=1.7, width=6.1, height=2.6, size=16)
bullets(s, [
    ("Prioritas segera:", 0, GREEN),
    ("Sesi ≥ 2 pakar forensik (blinded; κ ≥ 0,6) — formulir siap dari keluaran final: 4 indikator forensik + 18 benchmark", 1),
    ("Aktifkan kalibrasi & laporkan ECE/Brier/AURC setelah ≥ 4 label", 1),
    ("Kontrol negatif pada semua mode; user study H8", 1),
    ("Pengembangan:", 0, GREEN),
    ("satu jalur validasi; sensitivitas ambang aturan; NLI & reranker multibahasa; cakupan fakta SPO ke seluruh korpus; 50–100 jurnal per domain", 1),
], left=6.8, top=1.7, width=6.1, height=4.9, size=15)
callout(s, "Instrumen evaluasi pakar sudah lengkap dan teruji: formulir XLSX, kalkulator metrik + Cohen's κ, kalibrator otomatis aktif setelah ≥ 4 label",
        0.5, 4.6, 6.1, 1.5, fill=NAVY, size=14)
notes(s, "1:15", ["Katakan dengan tenang apa yang belum: evaluasi pakar. Lalu tunjukkan bahwa semua prasyaratnya sudah siap.",
                  "Saran disusun dari yang paling mendesak."],
      "Penutup.")

# 19. Terima kasih
s = prs.slides.add_slide(BLANK); _counter["n"] += 1
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
_txt(s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.8), Inches(1.2)), "Terima kasih", size=44, bold=True, color=WHITE)
_txt(s.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(11.8), Inches(1.6)),
     ["Sistem adalah alat bantu keputusan: mempersempit ruang pencarian dan menyediakan bukti yang dapat diaudit —",
      "penilaian induktif tetap pada peneliti."], size=20, color=RGBColor(0xE5, 0xE7, 0xEB))
_txt(s.shapes.add_textbox(Inches(0.8), Inches(5.4), Inches(11.8), Inches(1.2)),
     ["Kode & data: github.com/devnolife/llm-wizard-research", "894 unit test · 14 metrik · 2 korpus · 10 hipotesis dipetakan ke bukti"], size=16, color=RGBColor(0xCF, 0xD8, 0xE3))
notes(s, "0:15", ["Ucapkan terima kasih; siap menerima pertanyaan.", "Slide berikutnya adalah cadangan untuk Q&A — jangan ditampilkan kecuali ditanya."])

# ---------------------------------------------------------------- CADANGAN
s = new_slide("Slide Cadangan untuk Tanya-Jawab", section="CADANGAN")
bullets(s, [
    "C1 — Mengapa evaluasi pakar belum dilakukan, dan apa dampaknya?",
    "C2 — Mengapa dua LLM berbeda; bukankah itu mencampur variabel?",
    "C3 — Mengapa Rule Engine PASS 100%? Apakah ia berguna?",
    "C4 — Apa beda inkonsistensi vs heterogenitas; ketidaklengkapan vs ketiadaan dukungan bukti?",
    "C5 — Cacat yang ditemukan sendiri (dan diperbaiki) — bukti kejujuran metodologis",
    "C6 — Respons terhadap 8 kritik penguji pada proposal",
    "C7 — Ancaman validitas: kontaminasi pre-training, korpus kecil, non-determinisme",
    "C8 — Reproduksibilitas: berkas, seed, dan cara menjalankan ulang",
], size=18)
notes(s, "—", ["Peta cadangan; lompat ke slide sesuai pertanyaan."])

s = new_slide("C1 · Mengapa Evaluasi Pakar Belum Dilakukan?", section="CADANGAN")
bullets(s, [
    "Urutan yang benar: pakar menilai keluaran FINAL sistem, bukan versi tengah. Versi akhir (korroborasi, kopling dalam pipeline, perbaikan fusi verdict–kalibrasi) baru stabil September 2026",
    "Instrumen sudah lengkap dan teruji: formulir XLSX blinded (label genuine/trivial/illogical/already addressed, LCS, AS, peringkat, justifikasi REJECT), kalkulator EAR/FDR/LCS/AS/SHG/REP + Cohen's κ, kalibrator aktif otomatis ≥ 4 label",
    "Dampak dinyatakan eksplisit di BAB IV 4.4.5 & BAB V 5.1.3: keempat kriteria keberhasilan belum dapat dinyatakan terpenuhi maupun gagal; ECE/Brier/AURC belum dilaporkan",
    "Yang SUDAH terbukti tanpa pakar: deteksi dengan bukti terverifikasi, diskriminasi adversarial 6/6, signifikansi H9 & H10, false-gap 0,2/run, provenans 100%, stabilitas k/n",
    "Rencana: 2–3 pakar forensika digital, union beberapa run agar himpunan cukup besar, gold standard pakar sebelum melihat keluaran → recall",
], size=16)
notes(s, "—", ["Jawab tanpa defensif: ini keterbatasan waktu & akses, bukan desain. Tunjukkan prasyaratnya lengkap."])

s = new_slide("C2 · Mengapa Dua LLM Berbeda?", section="CADANGAN")
table(s, [
    ["Kebutuhan", "Ollama (llama3.2 / gpt-oss)", "Copilot SDK (claude-opus-4.8-fast)"],
    ["Seed & pengulangan untuk uji statistik", "ya — 5–7 run per mode, seed 43–49", "tidak dijamin deterministik"],
    ["Dokumen berbahasa Indonesia", "lemah pada model 3B", "memadai"],
    ["Biaya/kecepatan untuk ratusan run ablasi", "lokal, gratis", "kuota penyedia"],
    ["Dipakai untuk", "H6, H7, H9, H10, adversarial, kontrol negatif", "korpus aplikatif 35 jurnal (M9–M14)"],
], 0.5, 1.7, 12.3, col_w=[3.8, 4.2, 4.3], size=14, row_h=0.6)
bullets(s, [
    "Angka kedua kelompok TIDAK dibandingkan silang (BAB IV 4.2.5, keterbatasan 8); tiap kesimpulan sahih di dalam kelompoknya",
    "Sensitivitas model diamati langsung: model reasoning (gpt-oss) gagal pada JSON mode → mekanisme fallback",
    "Stabilitas pada model penyedia diukur empiris lewat k/n (3 run) dan 2 eksekusi versi akhir (3 dari 4 indikator muncul di keduanya)",
], top=5.0, height=1.9, size=15)
notes(s, "—", ["Intinya: pemisahan disengaja dan diumumkan, bukan kebetulan."])

s = new_slide("C3 · Rule Engine PASS 100% — Berguna?", section="CADANGAN")
bullets(s, [
    "PASS 100% pada data bersih adalah sifat INPUT (paper berkualitas tinggi), bukan bukti engine lemah — dibuktikan adversarial 6/6 (REJECT F1/F3, FLAG F2/K1/C1, PASS kontrol) dengan penurunan confidence bertingkat",
    "Multi-run: RERR full 3,8 ± 8,5% — pada satu dari lima run engine menandai indikator; PASS 100% bukan sifat tetap",
    "Kontribusi utama Rule Engine kualitatif: verdict menjadi bukti dalam kalibrasi (PASS ×1,10 / FLAG ×0,80 / REJECT batal), jejak aturan yang terpicu masuk provenans, dan ia menolak klaim adversarial yang lolos LLM",
    "Batas yang diakui (RQ3): engine hanya menalar atas fakta yang terekstrak; klaim tanpa entitas KG lolos secara default — karena itu cakupan fakta menjadi saran pengembangan",
    "H7 (FDR turun ≥ 20%) menunggu label pakar; proksi jumlah indikator tidak signifikan (power kecil, 5 run)",
], size=16)
notes(s, "—", ["Pertanyaan ini hampir pasti muncul. Kunci: adversarial + RERR multi-run + peran dalam kalibrasi."])

s = new_slide("C4 · Definisi yang Sering Dipertanyakan", section="CADANGAN")
table(s, [
    ["Pasangan", "Yang satu", "Yang lain", "Cara sistem membedakan"],
    ["Inkonsistensi vs heterogenitas", "temuan bertentangan pada klaim yang sebanding (arah efek berlawanan)", "hasil berbeda karena desain/populasi berbeda", "normalisasi klaim (arah, PICO) → NLI → Cochran Q, I², τ²; desain berbeda → bukan kontradiksi"],
    ["Ketidaklengkapan vs ketiadaan dukungan bukti", "aspek kritis TIDAK dibahas satu pun jurnal (sel kosong peta bukti)", "aspek DIBAHAS & DIKLAIM berulang, tetapi tanpa bukti primer yang dapat ditemukan", "peta cakupan aspek vs uji kegagalan retrieval leave-one-out (< 0,45) + citation echo"],
    ["Fragmentasi semantik vs struktural", "klaster embedding terpisah (isolasi 0,82)", "jurnal tidak berbagi rujukan & tidak saling mengutip (kopling 0,978)", "dua sinyal independen; struktural bebas LLM"],
    ["Korroborasi penulis vs gap", "keterbatasan/future work yang ditulis penulis = pernyataan satu penulis", "synthesis gap = klaim lintas-paper", "korroborasi hanya BUKTI untuk indikator; tidak pernah menjadi indikator atau skor"],
], 0.5, 1.6, 12.3, col_w=[2.5, 3.2, 3.2, 3.4], size=12, row_h=0.95)
notes(s, "—", ["Definisi Cooper/Booth (kritik #2) dan pembedaan tegas indikator 3 vs 4 (BAB III 3.6.1)."])

s = new_slide("C5 · Cacat yang Ditemukan Sendiri — dan Diperbaiki", section="CADANGAN")
table(s, [
    ["Cacat (silent, hanya terlihat pada korpus nyata)", "Akibat", "Perbaikan + uji regresi"],
    ["Asimetri istilah grounding vs kutipan (frasa panjang lolos grounding, tak pernah dikutip)", "semua indikator dipaksa needs_review (abstensi tidak diskriminatif)", "satu fungsi aspect_terms() untuk keduanya"],
    ["Indikator metodologis tanpa kutipan", "rantai provenans putus", "kalimat penyebut metode dijadikan bukti"],
    ["Peta bukti degeneratif 1×1 (100% terisi) tetap memancarkan indikator", "non-temuan disajikan sebagai gap", "gate ≥ 2×2 dan wajib ≥ 1 sel kosong"],
    ["Verdict final PASS tetapi kalibrasi memakai diskon FLAG ×0,80 (Rule Engine berjalan dua kali dengan pengait fakta berbeda)", "rekaman bertentangan dengan dirinya", "koordinator memfusikan ulang verdict final ke kalibrasi, abstensi, provenans; 4 uji"],
], 0.5, 1.6, 12.3, col_w=[5.2, 3.4, 3.7], size=13, row_h=0.85)
callout(s, "Pelajaran metodologis (BAB IV 4.4.3 butir 7): uji sintetis dengan data ideal tidak memadai untuk lapisan provenans; kini 35 + 4 uji regresi berbentuk data produksi; total 894 uji",
        0.5, 6.0, 12.3, 0.9, fill=NAVY, size=14)
notes(s, "—", ["Gunakan bila ditanya soal kualitas rekayasa atau 'apakah hasil bisa dipercaya'. Menemukan cacat sendiri = kredibilitas."])

s = new_slide("C6 · Respons terhadap 8 Kritik Penguji pada Proposal", section="CADANGAN")
table(s, [
    ["#", "Kritik", "Respons dalam tesis"],
    ["1", "Rumusan masalah bukan 'masalah'", "pola fenomena → kesenjangan → pertanyaan; 'Sejauh mana…' (BAB I 1.2)"],
    ["2", "Definisi synthesis gap tidak mainstream", "Cooper (1998) & Booth (2012): fragmentasi, inkonsistensi, ketidaklengkapan (+ dukungan bukti)"],
    ["3", "Batas kemampuan LLM tidak diakui", "BAB I 1.5 Batasan Epistemologis; RQ3; posisi decision support tool"],
    ["4", "RAG+LLM bukan hal baru", "Neuro-Symbolic Agentic: agen + Fact Table + Rule Engine + kalibrasi/provenans"],
    ["5", "Tak bisa bedakan semantik vs logis", "mekanisme 3 lapis + NLI; H9 terkonfirmasi"],
    ["6", "Perlu Rule-Based Validation Layer", "9 aturan / 3 verdict; adversarial 6/6; fusi ke kalibrasi"],
    ["7", "Diagram linear & black box", "4 fase + loop Observe–Think–Act–Evaluate; reasoning trace per indikator"],
    ["8", "KG tanpa Tabel Fakta SPO", "ontologi eksplisit: 8 entitas, 12 + 2 predikat; contoh transformasi teks → SPO (BAB III 3.5)"],
], 0.5, 1.6, 12.3, col_w=[0.5, 4.3, 7.5], size=13, row_h=0.55)
notes(s, "—", ["Penguji yang sama kemungkinan mengecek konsistensi revisi; slide ini menunjukkan setiap kritik punya jejak di bab mana."])

s = new_slide("C7 · Ancaman Validitas yang Diakui", section="CADANGAN")
bullets(s, [
    "Kontaminasi pre-training: paper benchmark (Transformer, BERT, ResNet) hampir pasti ada di data latih LLM → indikator bisa dari pengetahuan parametrik; mitigasi: grounding RAG, kutipan verbatim, validasi simbolis, dan korpus aplikatif lokal (jurnal Indonesia)",
    "Korpus kecil & indikator sedikit (2–4/topik): evaluasi statistik per korpus lemah → protokol multi-run dan union k/n; saran 50–100 jurnal/domain",
    "Non-determinisme: Jaccard antar-run 0,45–0,65; kritikus LLM stokastik (1–11 indikator akhir) → semua temuan LLM dilaporkan dengan k/n; angka satu run = realisasi",
    "Heterogenitas LLM antar rezim: tidak dibandingkan silang",
    "Kontrol negatif hanya pada cross-critic; 1 false gap berkeyakinan 0,885 lolos → provenans sebagai pertahanan terakhir",
    "Aspek parametrik ditandai: 1 dari 10 aspek 'tak dibahas' pada korpus forensik tidak ber-grounding korpus dan dilabeli parametrik, bukan disembunyikan",
], size=16)
notes(s, "—", ["Bila ditanya 'apa kelemahan terbesar?', jawab: evaluasi pakar tertunda dan stokastisitas — keduanya sudah diukur/dimitigasi, bukan diabaikan."])

s = new_slide("C8 · Reproduksibilitas", section="CADANGAN")
table(s, [
    ["Artefak", "Lokasi", "Isi"],
    ["Hasil benchmark & ablasi", "backend/experiments/results/experiment_<mode>_<model>.run<k>.json", "per run: seed, fakta, indikator, verdict, adversarial; multirun_stats_*.md (U, p, Holm, δ, CI)"],
    ["Korpus aplikatif", "data/processed/analysis_jobs.sqlite3 (job 977fdd5a → 5b2017ea)", "35 jurnal, 877 chunk, 4 indikator, provenans, korroborasi, kopling"],
    ["Penambangan gap 3 run", "data/processed/gaps_{d4eb6a1d,v3,v4}.jsonl → gaps_union*.jsonl", "528 gap union beranotasi run_support; novelty OpenAlex"],
    ["Instrumen pakar", "backend/experiments/expert_eval/", "expert_form_forensik.xlsx (4), expert_form_benchmark.xlsx (18), compute_metrics.py (EAR…κ)"],
    ["Uji & lint", "backend/tests (894 uji), flake8", "termasuk 35 uji regresi provenans + 4 uji fusi verdict–kalibrasi"],
    ["Menjalankan ulang", "POST /api/analysis-jobs/{id}/reanalyze; python -m experiments.run_multi", "job baru dari PDF yang sama dengan kode & LLM terkini"],
], 0.5, 1.6, 12.3, col_w=[2.6, 5.0, 4.7], size=12, row_h=0.72)
notes(s, "—", ["Bila ditanya 'bisa diulang?': ya — sebut seed, berkas, dan endpoint reanalyze."])

prs.save(OUT)
print(f"{OUT.name}: {len(prs.slides)} slide (19 inti + {len(prs.slides) - 19} cadangan)")
